from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\ofirm\OneDrive\Documents\ChatGPT\find Multi LLC owners")
OUT = ROOT / "Chicago" / "00_MARKET_ADAPTER" / "Chicago_MSA_Market_Adapter_and_Execution_Status.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, white=False, size=8):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(255, 255, 255) if white else RGBColor(0, 0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None, font_size=8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        shade(table.rows[0].cells[i], "000000")
        set_cell_text(table.rows[0].cells[i], header, bold=True, white=True, size=font_size)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return table


def field(run, name):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = name
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.7)
sec.right_margin = Inches(0.7)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(9)
styles["Title"].font.name = "Arial"
styles["Title"].font.size = Pt(24)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
for name, size in [("Heading 1", 16), ("Heading 2", 12), ("Heading 3", 10)]:
    styles[name].font.name = "Arial"
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("Chicago MSA Market Adapter  |  Research control document  |  Page ")
run.font.name = "Arial"
run.font.size = Pt(8)
field(footer.add_run(), "PAGE")
footer.add_run(" of ")
field(footer.add_run(), "NUMPAGES")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("CHICAGO MSA\nMARKET ADAPTER & EXECUTION STATUS")
r.bold = True
r.font.name = "Arial"
r.font.size = Pt(23)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("LLC portfolio identification by shared owner mailing address")
r.italic = True
r.font.size = Pt(12)
doc.add_paragraph()
add_table(doc, ["Control", "Value"], [
    ["Market", "Chicago-Naperville-Elgin, IL-IN Metropolitan Statistical Area (CBSA 16980)"],
    ["Boundary basis", "U.S. Census Bureau / OMB July 2023 delineation"],
    ["Governing specification", "00_MASTER_SPECIFICATION/LLC_Portfolio_Identification_and_Mailing_Master_Specification.docx"],
    ["Status date", "August 28, 2026"],
    ["Executed county adapters", "Cook County and Kendall County, Illinois"],
    ["Mailing status", "RESEARCH ONLY — no Chicago record is approved for printing yet"],
], [2.1, 5.7], 9)

doc.add_heading("1. Purpose and non-negotiable rule", level=1)
doc.add_paragraph(
    "This market adapter applies the generic master specification to Chicago and surrounding counties. "
    "It does not select multifamily parcels or require a unit-count threshold. Every parcel class remains in scope when the normalized legal owner name ends in LLC."
)
p = doc.add_paragraph()
r = p.add_run("The portfolio correlation key is the owner mailing address—not the property/situs address.")
r.bold = True
doc.add_paragraph(
    "A candidate portfolio exists only when the same normalized owner mailing address links at least two distinct LLC names and at least two distinct parcels. "
    "One outbound mail row is eventually produced per validated mailing address, after service-provider and postal suppressions."
)

doc.add_heading("2. Operating flow", level=1)
flow = [
    ["1", "Ingest all parcels", "Retain parcel ID, legal owner, owner mailing address, situs, source year and row lineage."],
    ["2", "Filter LLC owners", "Normalize the legal owner and require an LLC suffix at the end of the name."],
    ["3", "Normalize owner mailing address", "Standardize street, city, state and ZIP. Never substitute the situs address."],
    ["4", "Cluster", "Group different LLCs by the normalized owner mailing address."],
    ["5", "Qualify", "Require 2+ distinct LLCs and 2+ distinct parcel IDs."],
    ["6", "Suppress", "Remove lawyers, registered agents, CPAs/accountants, filing services, C/O and other service addresses."],
    ["7", "Postal validation", "Run CASS/DPV and campaign suppressions before print authorization."],
    ["8", "Mail", "Create one deduplicated postcard row per validated owner mailing address."],
]
add_table(doc, ["Step", "Stage", "Requirement"], flow, [0.55, 1.65, 5.6], 8)

doc.add_heading("3. Official market boundary", level=1)
doc.add_paragraph("The July 2023 federal delineation contains 13 counties. Kenosha County, Wisconsin is part of the wider combined statistical area, not this MSA, and is therefore excluded from this adapter.")
scope = [
    ["Cook", "IL", "17031", "Central"], ["DuPage", "IL", "17043", "Central"],
    ["Grundy", "IL", "17063", "Outlying"], ["McHenry", "IL", "17111", "Central"],
    ["Will", "IL", "17197", "Central"], ["DeKalb", "IL", "17037", "Outlying"],
    ["Kane", "IL", "17089", "Central"], ["Kendall", "IL", "17093", "Central"],
    ["Lake", "IL", "17097", "Central"], ["Jasper", "IN", "18073", "Outlying"],
    ["Lake", "IN", "18089", "Central"], ["Newton", "IN", "18111", "Outlying"],
    ["Porter", "IN", "18127", "Central"],
]
add_table(doc, ["County", "State", "FIPS", "Role"], scope, [2.2, 1.1, 1.5, 2.2], 8)

doc.add_heading("4. County access findings", level=1)
doc.add_paragraph(
    "County access is uneven. Production work may use official bulk downloads or stable APIs. Interactive sites are not automatically scraped when terms prohibit extraction or when bulk authorization is unclear. Those counties remain visible in the source matrix with a specific acquisition action."
)
findings = [
    ["Cook IL", "Executed", "Official Socrata API includes owner name, owner mailing address, situs and parcel ID. Publisher warns that owner/mailing fields are intermittently updated."],
    ["Kendall IL", "Executed", "Official live ArcGIS service exposes owner_name, owner_addr1/2, owner_csz, PIN and situs fields. Snapshot date is retained because the layer has no explicit tax year."],
    ["DuPage IL", "Verified lead", "Official ArcGIS parcel service exists; full real-estate database with billing fields is offered by purchase. Exact REST field coverage remains to be tested."],
    ["Lake IL", "Blocked for automated extraction", "The official tax-search site expressly prohibits mining/gathering/extraction. Use an authorized bulk request or FOIA route."],
    ["Kane IL", "Blocked pending acquisition", "County states that it has no open-data sharing policy and that GIS data can be acquired for a fee."],
    ["Will IL", "Partial", "Official parcel shapefile download is available; owner-mailing fields are not yet verified."],
    ["McHenry / DeKalb IL", "Access test needed", "Official interactive systems expose owner/assessment information; bulk endpoints and terms must be verified."],
    ["Grundy IL", "Discovery pending", "Official bulk owner-mail source has not yet been verified."],
    ["Porter IN", "Freshness blocked", "A public ArcGIS owner-mail layer exists but its name indicates a 2022 vintage, which is not current enough for a mailing campaign."],
    ["Lake / Jasper / Newton IN", "Discovery pending", "Official parcel/GIS routes exist; current bulk owner-mail schemas and terms remain to be verified."],
]
add_table(doc, ["County", "Status", "Finding"], findings, [1.5, 1.7, 4.8], 8)

doc.add_heading("5. Cook County adapter", level=1)
doc.add_paragraph("Official dataset: Cook County Assessor — Parcel Addresses (Socrata dataset 3723-97qp). The source was checked on August 28, 2026 and reported tax year 2026 as its latest year.")
fields = [
    ["pin", "parcel_id", "14-digit parcel identifier"],
    ["owner_address_name", "legal_owner_raw", "Owner mailing name; LLC-suffix filter source"],
    ["owner_address_full/city/state/zip", "owner mailing fields", "Exclusive portfolio grouping source"],
    ["prop_address_*", "property address fields", "Evidence and future mapping only"],
    ["row_id", "source_row_id", "Audit lineage"],
]
add_table(doc, ["Source field", "Canonical use", "Control"], fields, [2.4, 2.2, 3.4], 8)
doc.add_paragraph(
    "The first run downloaded 71,356 source rows matching a server-side *LLC suffix query. After local normalization and removal of one record with no usable owner mailing street, 71,355 parcel rows remained. Grouping them by normalized owner mailing address produced 3,505 clusters meeting both thresholds."
)
add_table(doc, ["Metric", "Result", "Interpretation"], [
    ["Source year", "2026", "Latest published year returned by the API"],
    ["Eligible LLC parcel rows", "71,355", "All property classes; one source row per parcel"],
    ["Qualifying address clusters", "3,505", "2+ distinct LLCs and 2+ distinct parcels"],
    ["Largest raw cluster", "70 LLCs", "High-collision addresses require heightened service-provider review"],
    ["Mail-ready rows", "0", "Service-provider and USPS gates have not yet been run"],
], [2.5, 1.5, 4.0], 8)

doc.add_heading("6. Kendall County adapter and partial-market merge", level=1)
doc.add_paragraph(
    "The official Kendall County live cadastral FeatureServer was queried on August 28, 2026. The service returned 5,045 owner names ending in LLC; 5,043 rows retained a usable owner mailing street. County-level grouping produced 110 qualifying address clusters."
)
doc.add_paragraph(
    "Cook and Kendall parcel evidence was then regrouped together by the same normalized owner mailing address. The partial MSA contains 3,627 qualifying clusters, including 25 addresses that link qualifying parcel evidence across both counties. This is partial coverage only; it is not represented as a completed 13-county market."
)
add_table(doc, ["Metric", "Result", "Control"], [
    ["Kendall LLC parcels", "5,043", "All parcel classes; live-source snapshot retained"],
    ["Kendall county clusters", "110", "2+ distinct LLCs and 2+ distinct parcels"],
    ["Cook + Kendall partial-MSA clusters", "3,627", "Re-grouped across completed county adapters"],
    ["Cross-county clusters", "25", "Same normalized owner-mail address appears in both counties"],
    ["Mail-ready rows", "0", "Professional/service-provider and USPS gates not complete"],
], [2.8, 1.5, 3.7], 8)

doc.add_heading("7. File map", level=1)
files = [
    ["00_MARKET_ADAPTER", "Chicago_MSA_Market_Adapter_and_Execution_Status.docx", "This control document"],
    ["01_FINAL_OUTPUTS", "cook_2026_candidate_mailing_addresses_RESEARCH_ONLY.csv", "Candidate address clusters; explicitly not printable"],
    ["02_SOURCE_DATA/00_BOUNDARY_REFERENCE", "list1_2023.xlsx", "Official Census/OMB county delineation"],
    ["03_PROCESSED", "cook_2026_llc_parcels.csv", "Parcel-level LLC evidence"],
    ["03_PROCESSED", "cook_2026_address_clusters.csv", "Address-level clusters and counts"],
    ["03_PROCESSED", "kendall_20260828_llc_parcels.csv", "Kendall parcel-level LLC evidence"],
    ["03_PROCESSED", "chicago_msa_partial_cook_kendall_20260828_address_clusters.csv", "Two-county normalized address clusters"],
    ["04_REPORTS_AND_WORKBOOKS", "Chicago_MSA_Source_Matrix_and_Data_Dictionary.xlsx", "Scope, sources, fields, output columns and status"],
    ["05_PIPELINE_CODE", "build_cook_llc_clusters.py", "Reproducible Cook API ingestion and grouping"],
    ["06_DOCUMENTATION", "README_Chicago_MSA.md", "Human-readable folder map and run instructions"],
    ["99_SUPPORT_FILES", "builders and QA renders", "Supporting scripts, previews and superseded documentation"],
]
add_table(doc, ["Folder", "File", "Purpose"], files, [2.2, 3.6, 2.2], 8)

doc.add_heading("8. Output columns and mailing gates", level=1)
doc.add_paragraph(
    "The parcel output preserves parcel_id, legal_owner_raw, legal_owner_normalized, owner mailing components, owner_mailing_address_key, situs components, source dataset ID and source row ID. The address-cluster output adds distinct_llc_count, property_count, LLC and parcel evidence lists, service_provider_review, usps_cass_dpv_status and mailing_status."
)
gates = [
    ["Legal owner", "PASS only when normalized name ends in LLC"],
    ["Cluster", "PASS only at 2+ distinct LLCs and 2+ distinct parcels"],
    ["Address provenance", "PASS only when the key comes from owner-mail fields"],
    ["Professional/service provider", "Suppress lawyer, registered-agent, CPA/accountant, filing-service and C/O destinations"],
    ["Collision/manual review", "Review high-count addresses and ambiguous suites before automation"],
    ["Postal", "CASS/DPV deliverability and standardized address required"],
    ["Deduplication", "Exactly one outbound row per standardized mailing address"],
    ["Print authorization", "MAIL_READY only when every required gate is PASS"],
]
add_table(doc, ["Gate", "Requirement"], gates, [2.2, 5.8], 8)

doc.add_heading("9. Next execution order", level=1)
for text in [
    "Complete Cook service-provider classification and evidence logging, starting with high-collision addresses.",
    "Run USPS CASS/DPV and campaign suppressions; create the first true Cook mail-ready file only after validation.",
    "Test DuPage REST fields and compare the authorized database-purchase option.",
    "Request authorized bulk extracts for counties whose sites prohibit or do not support automated extraction.",
    "Integrate each county into the same canonical schema, then deduplicate across county boundaries by standardized owner mailing address.",
]:
    doc.add_paragraph(text, style="List Number")

doc.add_heading("10. Primary references", level=1)
refs = [
    "U.S. Census Bureau, July 2023 delineation files: https://www.census.gov/geographies/reference-files/time-series/demo/metro-micro/delineation-files.html",
    "Cook County Assessor, Parcel Addresses: https://datacatalog.cookcountyil.gov/d/3723-97qp",
    "Cook County 2025 Open Data Refresh: https://datacatalog.cookcountyil.gov/stories/s/Assessor-2025-Open-Data-Refresh/gzdr-q7c4/",
    "DuPage County parcel service: https://gis.dupageco.org/arcgis/rest/services/DuPage_County_IL/ParcelsWithRealEstateCC/MapServer",
    "Lake County, Illinois property search terms: https://tax.lakecountyil.gov/search/commonsearch.aspx?mode=realprop",
    "Kane County GIS FAQ: https://www.kanecountyil.gov/Pages/gis-faq.aspx",
    "Will County vector downloads: https://willcounty.gov/County-Offices/Administration/GIS-Division/Data/Vector",
    "Kendall County live cadastral service: https://maps.co.kendall.il.us/server/rest/services/Hosted/Current_Cadastral_Features/FeatureServer/1",
]
for ref in refs:
    doc.add_paragraph(ref, style="List Bullet")

doc.core_properties.title = "Chicago MSA Market Adapter and Execution Status"
doc.core_properties.subject = "LLC portfolio identification using shared owner mailing addresses"
doc.core_properties.author = "Research Operations"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
